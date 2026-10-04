#!/usr/bin/env python3
"""Builds the committed workbook fixtures.

- *.xlsx through openpyxl (a real, widely used writer: shared strings, styles, dates,
  the 1904 epoch, hidden sheets, rich text).
- *.ods by hand, straight from the ODF 1.3 schema (no ODS writer is installed here).

Re-run to regenerate; the outputs are committed so the test suite needs no Python.
"""

import datetime as dt
import os
import zipfile

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.utils.datetime import CALENDAR_MAC_1904

HERE = os.path.dirname(os.path.abspath(__file__))


def basic(epoch=None):
    wb = Workbook()
    if epoch is not None:
        wb.epoch = epoch
    ws = wb.active
    ws.title = "Data"
    ws.append(["id", "name", "amount", "when", "flag", "note", "ratio", "day", "at"])
    ws.append([1, "alice", 2.5, dt.datetime(2024, 1, 31, 10, 30, 0), True, "plain", 0.25, dt.date(2023, 3, 15), dt.time(15, 4, 5)])
    ws.append([2, "bob, jr", -7, dt.datetime(2024, 2, 29, 0, 0, 0), False, None, 1e-3, dt.date(1900, 2, 28), dt.time(0, 0, 0)])
    ws.append([3, "", 1e15, None, None, "x <&> \"y\"", 100, dt.date(9999, 12, 31), dt.time(23, 59, 59)])
    # Row 5 is left absent: a gap in numbering.
    ws.cell(row=6, column=1, value=6)
    ws.cell(row=6, column=26, value="Z6")  # sparse: columns B..Y empty
    ws.cell(row=7, column=1, value="=A2*2")  # formula, no cached value from openpyxl
    ws["C1"].number_format = "0.00"
    for row in range(2, 4):
        ws.cell(row=row, column=4).number_format = "yyyy-mm-dd hh:mm:ss"
        ws.cell(row=row, column=8).number_format = "yyyy-mm-dd"
        ws.cell(row=row, column=9).number_format = "hh:mm:ss"
        ws.cell(row=row, column=7).number_format = "0.00%"
    ws.cell(row=4, column=8).number_format = "yyyy-mm-dd"
    ws.cell(row=4, column=9).number_format = "hh:mm:ss"
    return wb


def multisheet():
    wb = Workbook()
    first = wb.active
    first.title = "First"
    first.append(["a"])
    first.append([1])
    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["h"])
    hidden.append([2])
    very = wb.create_sheet("Very")
    very.sheet_state = "veryHidden"
    very.append(["v"])
    very.append([3])
    last = wb.create_sheet("Last")
    last.append(["l"])
    last.append([4])
    return wb


def rich():
    wb = Workbook()
    ws = wb.active
    ws.title = "Rich"
    ws.append(["text"])
    ws["A2"] = CellRichText("plain ", TextBlock(InlineFont(b=True), "bold"), " tail")
    ws["A3"] = "shared once"
    ws["A4"] = "shared once"
    return wb


def ods_basic():
    content = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:calcext="urn:org:documentfoundation:names:experimental:calc:xmlns:calcext:1.0" xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" office:version="1.3">
<office:body><office:spreadsheet>
<table:table table:name="Data">
<table:table-column table:number-columns-repeated="1024"/>
<table:table-header-rows>
<table:table-row>
<table:table-cell office:value-type="string"><text:p>id</text:p></table:table-cell>
<table:table-cell office:value-type="string"><text:p>amount</text:p></table:table-cell>
<table:table-cell office:value-type="string"><text:p>pct</text:p></table:table-cell>
<table:table-cell office:value-type="string"><text:p>money</text:p></table:table-cell>
<table:table-cell office:value-type="string"><text:p>day</text:p></table:table-cell>
<table:table-cell office:value-type="string"><text:p>stamp</text:p></table:table-cell>
<table:table-cell office:value-type="string"><text:p>elapsed</text:p></table:table-cell>
<table:table-cell office:value-type="string"><text:p>flag</text:p></table:table-cell>
<table:table-cell office:value-type="string"><text:p>text</text:p></table:table-cell>
<table:table-cell office:value-type="string"><text:p>err</text:p></table:table-cell>
</table:table-row>
</table:table-header-rows>
<table:table-row>
<table:table-cell office:value-type="float" office:value="1"><text:p>1</text:p></table:table-cell>
<table:table-cell office:value-type="float" office:value="2.5"><text:p>2.5</text:p></table:table-cell>
<table:table-cell office:value-type="percentage" office:value="0.25"><text:p>25%</text:p></table:table-cell>
<table:table-cell office:value-type="currency" office:currency="USD" office:value="-7"><text:p>-$7.00</text:p></table:table-cell>
<table:table-cell office:value-type="date" office:date-value="2023-03-15"><text:p>2023-03-15</text:p></table:table-cell>
<table:table-cell office:value-type="date" office:date-value="2024-01-31T10:30:00"><text:p>x</text:p></table:table-cell>
<table:table-cell office:value-type="time" office:time-value="PT25H30M00S"><text:p>25:30:00</text:p></table:table-cell>
<table:table-cell office:value-type="boolean" office:boolean-value="true"><text:p>TRUE</text:p></table:table-cell>
<table:table-cell office:value-type="string"><text:p>a<text:s text:c="3"/>b<text:tab/>c</text:p><text:p>second &amp; &lt;line&gt;</text:p><office:annotation><text:p>ignored</text:p></office:annotation></table:table-cell>
<table:table-cell table:formula="of:=1/0" office:value-type="string" calcext:value-type="error"><text:p>#DIV/0!</text:p></table:table-cell>
</table:table-row>
<table:table-row table:number-rows-repeated="2">
<table:table-cell office:value-type="float" office:value="9"><text:p>9</text:p></table:table-cell>
<table:table-cell table:number-columns-repeated="3"/>
<table:table-cell office:value-type="string" office:string-value="sv"><text:p>ignored body</text:p></table:table-cell>
<table:table-cell office:value-type="float" office:value="4" table:number-columns-repeated="2"><text:p>4</text:p></table:table-cell>
</table:table-row>
<table:table-row table:number-rows-repeated="1048570">
<table:table-cell table:number-columns-repeated="1024"/>
</table:table-row>
</table:table>
<table:table table:name="Second">
<table:table-row><table:table-cell office:value-type="string"><text:p>only</text:p></table:table-cell></table:table-row>
<table:table-row><table:table-cell office:value-type="float" office:value="42"><text:p>42</text:p></table:table-cell></table:table-row>
</table:table>
</office:spreadsheet></office:body></office:document-content>
"""
    manifest = """<?xml version="1.0" encoding="UTF-8"?>
<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.3">
<manifest:file-entry manifest:full-path="/" manifest:version="1.3" manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/>
<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>
</manifest:manifest>
"""
    path = os.path.join(HERE, "basic.ods")
    with zipfile.ZipFile(path, "w") as zf:
        # The mimetype entry must be first and stored (ODF 1.3 Part 2 §3.3).
        zf.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.spreadsheet", compress_type=zipfile.ZIP_STORED)
        zf.writestr("META-INF/manifest.xml", manifest, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("content.xml", content, compress_type=zipfile.ZIP_DEFLATED)


def main():
    basic().save(os.path.join(HERE, "basic.xlsx"))
    basic(CALENDAR_MAC_1904).save(os.path.join(HERE, "basic-1904.xlsx"))
    multisheet().save(os.path.join(HERE, "multisheet.xlsx"))
    rich().save(os.path.join(HERE, "rich.xlsx"))
    ods_basic()


if __name__ == "__main__":
    main()
